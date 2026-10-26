# -*- coding: utf-8 -*-
"""
importar.py  ·  Trae movimientos de golpe
=========================================
Tres caminos, todos con vista previa antes de guardar nada:
  1. El CSV «Plusvalías y minusvalías» de cada fondo de MyInvestor.
  2. La plantilla genérica (Excel o CSV), rellena a mano o por una IA.
  3. El texto que devuelve la IA, pegado tal cual.

El proceso tiene dos pasos: preparar() lo lee, reconoce los productos y deja un
plan; aplicar() ejecuta ese plan sobre la cartera. La vista previa es aplicar()
sobre una copia, así que lo que ves es exactamente lo que se guardará.
"""

import copy
import csv
import datetime as dt
import io
import re
import unicodedata

from . import almacen, buscar, motor
from .motor import num_es, valor_en

COLUMNAS = ["fecha", "identificador", "nombre", "tipo_producto", "tipo_movimiento",
            "unidades", "importe", "moneda", "comision", "cartera", "nota"]

# Nombres alternativos que se aceptan en la cabecera.
SINONIMOS = {
    "fecha": ["fecha", "date", "fecha operacion", "fecha valor", "fecha ejecucion"],
    "identificador": ["identificador", "isin", "ticker", "simbolo", "codigo", "id"],
    "nombre": ["nombre", "producto", "descripcion", "name"],
    "tipo_producto": ["tipo producto", "tipo_producto", "clase", "tipo de producto"],
    "tipo_movimiento": ["tipo movimiento", "tipo_movimiento", "movimiento", "operacion", "tipo"],
    "unidades": ["unidades", "participaciones", "titulos", "acciones", "cantidad", "units"],
    "importe": ["importe", "importe eur", "total", "amount", "importe total"],
    "moneda": ["moneda", "divisa", "currency"],
    "comision": ["comision", "comisiones", "fee", "gastos"],
    "cartera": ["cartera", "portfolio", "cartera nombre"],
    "nota": ["nota", "notas", "comentario", "observaciones"],
}
TIPO_MOV = {
    "compra": "compra", "suscripcion": "compra", "buy": "compra", "aportacion": "compra",
    "venta": "venta", "reembolso": "venta", "sell": "venta",
    "dividendo": "dividendo", "cupon": "dividendo", "dividend": "dividendo", "interes": "dividendo",
    "comision": "comision", "fee": "comision", "custodia": "comision",
    "saldo": "saldo", "valor": "saldo", "valoracion": "saldo",
    "traspaso salida": "traspaso_salida", "traspaso de salida": "traspaso_salida", "traspaso origen": "traspaso_salida",
    "traspaso entrada": "traspaso_entrada", "traspaso de entrada": "traspaso_entrada", "traspaso destino": "traspaso_entrada",
    "ajuste": "ajuste", "ajuste de participaciones": "ajuste",
}
TIPO_PROD = {
    "fondo": "fondo", "fondo de inversion": "fondo", "etf": "etf", "etp": "etf",
    "accion": "accion", "acciones": "accion", "cripto": "cripto", "criptomoneda": "cripto",
    "commodity": "commodity", "materia prima": "commodity", "oro": "commodity",
    "bono": "bono", "renta fija": "bono", "pension": "pension", "plan de pensiones": "pension",
    "efectivo": "efectivo", "cuenta": "efectivo", "cuenta corriente": "efectivo",
    "inmueble": "inmueble", "piso": "inmueble", "vivienda": "inmueble",
    "deuda": "deuda", "prestamo": "deuda", "hipoteca": "deuda", "otro": "otro",
    "p2p": "p2p", "prestamos p2p": "p2p", "crowdfunding": "p2p", "crowdlending": "p2p",
}


def sin_tildes(t):
    t = unicodedata.normalize("NFKD", str(t or "")).encode("ascii", "ignore").decode()
    return re.sub(r"\s+", " ", t.replace("_", " ")).strip().lower()


def decodifica(crudo):
    """Bytes de un archivo de texto -> str, probando las codificaciones habituales."""
    if isinstance(crudo, str):
        return crudo
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return crudo.decode(enc)
        except UnicodeDecodeError:
            continue
    return crudo.decode("utf-8", "replace")


# ---------------------------------------------------------------- MyInvestor

CAB_MYINVESTOR = {"fecha": ["fecha fiscal", "fecha"],
                  "coste": ["inversion", "inversión", "coste"],
                  "valor": ["valor de mercado", "valor mercado", "valor"]}


def leer_csv_myinvestor(texto):
    """
    Lee el CSV "Plusvalías y minusvalías" de un fondo de MyInvestor.
    Devuelve (lotes, reembolsos): lotes = [[fecha, coste, valor]] de lo que sigues
    teniendo; reembolsos = [[fecha, resultado]] de lo ya vendido, del que el extracto
    solo da la plusvalía.
    """
    texto = decodifica(texto)
    delim = ";" if texto.count(";") >= texto.count(",") else ","
    filas = [f for f in csv.reader(io.StringIO(texto), delimiter=delim)
             if any((c or "").strip() for c in f)]
    if not filas:
        return [], []

    cab = [(c or "").strip().lower() for c in filas[0]]

    def col(clave, defecto):
        for i, c in enumerate(cab):
            if any(a in c for a in CAB_MYINVESTOR[clave]):
                return i
        return defecto

    ic, ico, iv = col("fecha", 0), col("coste", 1), col("valor", 2)
    lotes, reembolsos = [], []
    for fila in filas[1:]:
        if len(fila) <= max(ic, ico, iv):
            continue
        m = re.search(r"(\d{1,2})[/-](\d{1,2})[/-](\d{4})", fila[ic])
        if not m:
            continue
        fecha = f"{m.group(3)}-{int(m.group(2)):02d}-{int(m.group(1)):02d}"
        coste, valor = num_es(fila[ico]), num_es(fila[iv])
        if coste <= 0 and valor <= 0:
            # Participaciones ya reembolsadas: solo queda el resultado fiscal.
            if len(fila) > 3 and num_es(fila[3]):
                reembolsos.append([fecha, round(num_es(fila[3]), 2)])
            continue
        lotes.append([fecha, round(coste, 2), round(valor, 2)])
    lotes.sort(key=lambda x: x[0])
    return lotes, reembolsos


# Diferencia a partir de la cual un lote no es una compra normal (sus dos cálculos de
# participaciones no casan): viene de un traspaso. En compras normales casan a ~0,3 %.
TOLERANCIA_TRASPASO = 0.02


def vl_del_extracto(lotes, serie_vl):
    """
    Valor liquidativo con el que MyInvestor valoró el extracto (suele ser el de uno o
    dos días antes de descargarlo). Es el de la serie que hace cuadrar más lotes:
    en una compra normal, coste / VL del día de compra = valor de mercado / ese VL.
    A igualdad, el que deja menos error. Devuelve (vl, lotes que cuadran); si no
    cuadra ninguno, el último VL y 0.
    """
    normales = [(c / serie_vl[f], v) for f, c, v in lotes if serie_vl.get(f) and c > 0 and v > 0]
    if not serie_vl:
        return None, 0
    mejor, nota = None, (0, 0.0)
    for f in sorted(serie_vl, reverse=True):
        vl = serie_vl[f]
        errores = [abs(u * vl / v - 1) for u, v in normales]
        dentro = [e for e in errores if e < 0.005]
        candidata = (len(dentro), -sum(dentro) / len(dentro) if dentro else 0.0)
        if candidata > nota:
            mejor, nota = vl, candidata
    return (mejor, nota[0]) if nota[0] else (serie_vl[max(serie_vl)], 0)


def myinvestor_a_movimientos(producto_id, lotes, reembolsos, serie_vl):
    """
    Convierte los lotes en compras. El extracto no trae participaciones, así que
    se calculan como coste / valor liquidativo del día de compra, que es exacto.

    Excepción, los traspasos: el lote conserva la fecha y el coste del fondo de origen
    (así lo exige Hacienda), y dividir ese coste entre el VL de este fondo en esa fecha
    da participaciones falsas. Se detectan porque no casan con el valor de mercado, y
    sus participaciones se sacan de valor de mercado / VL del extracto.

    Devuelve (movimientos, lotes sin VL de ese día, lotes de traspaso).
    """
    vl_extracto, cuadran = vl_del_extracto(lotes, serie_vl)
    movs, sin_vl, traspasos = [], 0, 0
    for fecha, coste, valor in lotes:
        vc = serie_vl.get(fecha)
        nota = "MyInvestor"
        es_traspaso = False
        if vc and coste > 0:
            unidades = coste / vc
            if cuadran and valor > 0 and abs(unidades * vl_extracto / valor - 1) > TOLERANCIA_TRASPASO:
                unidades, traspasos = valor / vl_extracto, traspasos + 1
                nota = "MyInvestor · traspaso: participaciones estimadas con el valor de mercado"
                es_traspaso = True
        elif vl_extracto and valor > 0:
            unidades = valor / vl_extracto
            inicio = min(serie_vl) if serie_vl else None
            if inicio and fecha < (dt.date.fromisoformat(inicio) - dt.timedelta(days=7)).isoformat():
                # Anterior a toda la historia del fondo: no se pudo comprar aquí, viene de un traspaso.
                traspasos, es_traspaso = traspasos + 1, True
                nota = "MyInvestor · traspaso: participaciones estimadas con el valor de mercado"
            else:
                sin_vl += 1
        else:
            unidades, sin_vl = 0.0, sin_vl + 1
        movs.append({"fecha": fecha, "producto": producto_id, "tipo": "compra",
                     "unidades": round(unidades, 6), "importe": coste, "nota": nota,
                     **({"_traspaso": True} if es_traspaso else {})})
    for fecha, resultado in reembolsos:
        # Del reembolso solo se conoce la plusvalía: se anota como una venta de 0
        # participaciones que cobra ese resultado.
        movs.append({"fecha": fecha, "producto": producto_id, "tipo": "venta",
                     "unidades": 0, "importe": resultado,
                     "nota": "Plusvalía de un reembolso (MyInvestor no da la fecha de venta)"})
    return movs, sin_vl, traspasos


# ---------------------------------------------------------------- lectura de tablas

def leer_tabla(nombre, contenido):
    """Excel (.xlsx), CSV o texto pegado -> [(número de fila, {columna: valor})]."""
    if nombre and nombre.lower().endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook
        libro = load_workbook(io.BytesIO(contenido), data_only=True, read_only=True)
        hoja = libro["Movimientos"] if "Movimientos" in libro.sheetnames else libro.worksheets[0]
        filas = [list(f) for f in hoja.iter_rows(values_only=True)]
    else:
        texto = decodifica(contenido)
        # Lo que devuelve una IA suele venir dentro de un bloque ``` ... ```.
        texto = re.sub(r"^\s*```[a-zA-Z]*\s*$", "", texto, flags=re.M).strip()
        lineas = texto.splitlines()
        if not any(l.strip() for l in lineas):
            return [], "El texto está vacío."
        cab = next(l for l in lineas if l.strip())
        delim = max([";", "\t", ","], key=cab.count)
        filas = list(csv.reader(io.StringIO("\n".join(lineas)), delimiter=delim))
    # Se guarda el número de fila real (el que ves en Excel) aunque haya filas en blanco.
    filas = [(n, f) for n, f in enumerate(filas, start=1)
             if f and any(c not in (None, "") and str(c).strip() for c in f)]
    if not filas:
        return [], "No hay ninguna fila con datos."

    cab = [sin_tildes(c) for c in filas[0][1]]
    idx = {}
    for col, alias in SINONIMOS.items():
        for i, c in enumerate(cab):
            if c in alias and i not in idx.values():
                idx[col] = i
                break
    faltan = [c for c in ("fecha", "tipo_movimiento", "importe") if c not in idx]
    if faltan:
        return [], ("No encuentro las columnas " + ", ".join(faltan) + ". La primera fila tiene que ser la "
                    "cabecera de la plantilla: " + ";".join(COLUMNAS))
    if len(filas) < 2:
        if nombre and nombre.lower().endswith((".xlsx", ".xlsm")):
            return [], ("La hoja «Movimientos» está vacía: solo tiene la cabecera. Las filas de la hoja "
                        "«Ejemplo» son de muestra y no se importan; escribe tus operaciones en «Movimientos», "
                        "debajo de la cabecera, guarda el archivo y vuelve a subirlo.")
        return [], "Solo hay cabecera: escribe tus operaciones debajo, una por fila."
    salida = []
    for n, f in filas[1:]:
        salida.append((n, {col: (f[i] if i < len(f) else None) for col, i in idx.items()}))
    return salida, None


def lee_fecha(v):
    if isinstance(v, dt.datetime):
        return v.date().isoformat()
    if isinstance(v, dt.date):
        return v.isoformat()
    t = str(v or "").strip()
    m = re.match(r"^(\d{4})-(\d{1,2})-(\d{1,2})", t)
    if m:
        a, me, di = int(m.group(1)), int(m.group(2)), int(m.group(3))
    else:
        m = re.match(r"^(\d{1,2})[/.-](\d{1,2})[/.-](\d{2,4})", t)
        if not m:
            return None
        a = int(m.group(3)) + (2000 if len(m.group(3)) == 2 else 0)
        me, di = int(m.group(2)), int(m.group(1))
    try:
        return dt.date(a, me, di).isoformat()
    except ValueError:   # un 31 de febrero, un mes 13...
        return None


def lee_numero(v):
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    if isinstance(v, (int, float)):
        return abs(float(v))
    t = str(v).strip().replace("€", "").replace("$", "").replace(" ", "")
    if not re.fullmatch(r"[-+]?[\d.,]+", t):
        return "error"
    return abs(num_es(t))


# ---------------------------------------------------------------- preparar el plan

class Plan:
    """Lo que se va a importar: productos nuevos, movimientos, saldos y errores."""

    def __init__(self):
        self.productos_nuevos = []   # [{"ref": "nuevo:1", "datos": {...}, "precio": ..., "fecha": ...}]
        self.movimientos = []        # [{"fila", "producto" (id o ref), campos del movimiento, "marcas"}]
        self.valoraciones = []
        self.reemplazar = []         # [(producto id o ref, origen)]: se borran sus importados previos
        self.errores = []            # [{"fila", "mensaje"}]
        self.avisos = []
        # Fondos de MyInvestor con lotes que vienen de un traspaso: se pregunta de qué fondo venían.
        self.traspasos_mi = []       # [{"ref", "isin", "lotes", "coste", "unidades", "serie", "sugerencia", "detalle"}]
        self.series_mi = {}          # {ref: valor liquidativo diario} de cada fondo de MyInvestor importado

    def error(self, fila, mensaje):
        self.errores.append({"fila": fila, "mensaje": mensaje})


def _coincidencias(cfg, identificador, nombre):
    """Todos los productos que casan con ese identificador o nombre (puede haber
    varios si el mismo ISIN está en distintas carteras)."""
    ident = (identificador or "").strip().upper()
    if ident:
        encontrados = [p for p in cfg.get("productos", [])
                       if ident in ((p.get("identificador") or "").upper(), (p.get("codigo") or "").upper())]
        if encontrados:
            return encontrados
    nom = sin_tildes(nombre)
    if nom and not ident:
        return [p for p in cfg.get("productos", [])
                if nom in (sin_tildes(p.get("nombre")), sin_tildes(p.get("corto")))]
    return []


def _producto_existente(cfg, identificador, nombre, cartera=""):
    candidatos = _coincidencias(cfg, identificador, nombre)
    if not candidatos:
        return None
    cart = (cartera or "").strip().lower()
    if cart:
        for p in candidatos:
            if (p.get("cartera") or "").strip().lower() == cart:
                return p
        return None
    return candidatos[0]


class Resolutor:
    """Encuentra (o propone crear) el producto de cada fila, una sola vez por producto."""

    def __init__(self, cfg, plan, carpeta):
        self.cfg, self.plan, self.carpeta = cfg, plan, carpeta
        self.cache, self.series, self.frescas = {}, {}, set()

    def producto(self, fila, identificador, nombre, tipo_prod, cartera=""):
        cart = (cartera or "").strip()
        clave = ((identificador or "").strip().upper(), sin_tildes(nombre) if not identificador else "",
                 cart.lower())
        if clave in self.cache:
            return self.cache[clave]
        ref = None
        candidatos = _coincidencias(self.cfg, identificador, nombre)
        if len(candidatos) > 1 and not cart and ((identificador or "").strip()):
            # Mismo ISIN en varios productos/carteras y la fila no dice a cuál va:
            # no se adivina, se pide que lo indique en la columna «cartera».
            nombres = ", ".join(f"«{(p.get('corto') or p.get('nombre'))}»"
                                + (f" (cartera {p.get('cartera')})" if p.get("cartera") else " (sin cartera)")
                                for p in candidatos)
            self.plan.error(fila, f"«{(identificador or '').strip()}» está en varios productos: {nombres}. "
                                  "Pon en la columna «cartera» a cuál va cada fila (p. ej. Metal) y vuelve a importar. "
                                  "Si es para una cartera nueva, pon su nombre y se creará una copia separada.")
            self.cache[clave] = None
            return None
        existente = _producto_existente(self.cfg, identificador, nombre, cart)
        if existente:
            ref = existente["id"]
        elif identificador:
            n_fallos = len(buscar.FALLOS)
            res = buscar.buscar(identificador)
            if not res and buscar.hubo_fallos_desde(n_fallos):
                self.plan.error(fila, f"No he podido consultar «{identificador}» en internet ahora mismo "
                                      "(la conexión ha fallado o Yahoo/Morningstar no han respondido). "
                                      "Comprueba tu conexión y vuelve a pulsar «Previsualizar»: suele bastar "
                                      "con intentarlo otra vez.")
                ref = None
            elif res:
                r = res[0]
                fi = r.get("ficha") or {}
                datos = {"nombre": nombre or r.get("nombre") or r["codigo"], "tipo": r["tipo"],
                         "identificador": identificador.strip(), "fuente": r["fuente"], "codigo": r["codigo"],
                         "moneda": r.get("moneda") or "EUR", "vivo": r.get("vivo") or "",
                         "ter": fi.get("ter"), "riesgo": fi.get("riesgo"), "clase": fi.get("clase") or "",
                         "cartera": cart, "gestora": fi.get("gestora") or "", "largoPlazo": True}
                ref = self._nuevo(datos, r)
            else:
                self.plan.error(fila, f"No encuentro «{identificador}» con precio en internet. Si no cotiza, "
                                      "créalo en Productos con «Valor anotado a mano» y ese mismo "
                                      "identificador, y vuelve a importar.")
        elif nombre:
            tipo = tipo_prod if tipo_prod in motor.TIPOS else "otro"
            ref = self._nuevo({"nombre": nombre, "tipo": tipo, "fuente": "manual", "cartera": cart,
                               "largoPlazo": tipo not in ("efectivo", "deuda")}, None)
        else:
            self.plan.error(fila, "No tiene identificador ni nombre: no sé de qué producto es.")
        self.cache[clave] = ref
        return ref

    def _nuevo(self, datos, resultado):
        ref = f"nuevo:{len(self.plan.productos_nuevos) + 1}"
        self.plan.productos_nuevos.append({
            "ref": ref, "datos": datos,
            "precio": resultado and resultado.get("precio"), "monedaPrecio": resultado and resultado.get("moneda"),
            "fecha": resultado and resultado.get("fecha"), "mercado": resultado and resultado.get("mercado")})
        return ref

    def datos(self, ref):
        if ref and ref.startswith("nuevo:"):
            return next(n["datos"] for n in self.plan.productos_nuevos if n["ref"] == ref)
        return almacen.producto(self.cfg, ref)

    def serie(self, ref, fecha=None):
        """Precios en euros del producto; si no llegan a 'fecha', se bajan de nuevo."""
        if ref not in self.series:
            p = self.datos(ref)
            self.series[ref] = motor.serie_producto(p, self.carpeta) if p and p.get("fuente") != "manual" else {}
        s = self.series[ref]
        if fecha and s and max(s) < fecha and ref not in self.frescas:
            self.frescas.add(ref)
            self.series[ref] = s = motor.serie_producto(self.datos(ref), self.carpeta, fresca=True)
        return s


def preparar_tabla(cfg, filas, carpeta):
    """Plantilla o texto de la IA -> Plan."""
    plan = Plan()
    res = Resolutor(cfg, plan, carpeta)
    cambios, traspasos = {}, []
    for n, f in filas:
        fecha = lee_fecha(f.get("fecha"))
        tipo = TIPO_MOV.get(sin_tildes(f.get("tipo_movimiento")))
        ident = str(f.get("identificador") or "").strip()
        nombre = str(f.get("nombre") or "").strip()
        tprod = TIPO_PROD.get(sin_tildes(f.get("tipo_producto")), "")
        importe, unidades, comision = (lee_numero(f.get(k)) for k in ("importe", "unidades", "comision"))
        moneda = (str(f.get("moneda") or "").strip() or "EUR").upper()
        cartera = str(f.get("cartera") or "").strip()
        nota = str(f.get("nota") or "").strip()

        if not fecha:
            plan.error(n, f"La fecha «{f.get('fecha') or ''}» no es válida. Usa el formato 2025-03-10 o 10/03/2025.")
            continue
        if not tipo:
            plan.error(n, f"El tipo de movimiento «{f.get('tipo_movimiento') or ''}» no es válido: "
                          "usa compra, venta, dividendo, comision, saldo, traspaso_salida o traspaso_entrada.")
            continue
        if "error" in (importe, unidades, comision):
            plan.error(n, "Algún número no se entiende (importe, unidades o comisión).")
            continue
        if not importe and tipo not in ("saldo", "ajuste"):
            plan.error(n, "No tiene importe.")
            continue
        if tipo == "ajuste":
            importe = importe or 0.0
        if not ident and not nombre:
            plan.error(n, "No tiene identificador: pon el ISIN o el ticker (o, si es una cuenta, su nombre).")
            continue
        if tipo == "saldo" and not ident:
            tprod = tprod or "efectivo"
        ref = res.producto(n, ident, nombre, tprod, cartera)
        if not ref:
            continue
        marcas = []

        # Importes en otra moneda: se pasan a euros con el cambio de ese día.
        if moneda != "EUR" and importe is not None:
            if moneda not in cambios:
                cambios[moneda] = motor.serie_cambio(moneda, carpeta)
            fx = valor_en(cambios[moneda], fecha, margen=6)
            if not fx and (moneda, "fresco") not in cambios:
                # El cambio guardado no llega a esa fecha: se baja de nuevo una vez.
                cambios[moneda] = motor.serie_cambio(moneda, carpeta, fresca=True)
                cambios[(moneda, "fresco")] = True
                fx = valor_en(cambios[moneda], fecha, margen=6)
            if not fx:
                plan.error(n, f"No encuentro el cambio de {moneda} a euros del {almacen.fmt_fecha(fecha)}.")
                continue
            importe = importe * fx
            comision = comision * fx if comision else comision
            marcas.append(f"convertido de {moneda}")

        if tipo == "saldo":
            plan.valoraciones.append({"fila": n, "producto": ref, "fecha": fecha, "valor": round(importe or 0, 2)})
            continue

        p = res.datos(ref)
        cotiza = p and p.get("fuente") != "manual" and p.get("tipo") not in almacen.SOLO_SALDO
        if tipo in ("traspaso_salida", "traspaso_entrada"):
            if not unidades and cotiza:
                precio = valor_en(res.serie(ref, fecha), fecha, margen=6)
                if not precio:
                    plan.error(n, "No trae participaciones y no encuentro el precio de ese día para calcularlas.")
                    continue
                unidades = importe / precio
                marcas.append("unidades calculadas")
            traspasos.append({"fila": n, "producto": ref, "fecha": fecha, "tipo": tipo, "unidades": unidades or 0,
                              "importe": round(importe, 2), "nota": nota, "marcas": marcas})
            continue
        if tipo in ("compra", "venta") and not unidades and cotiza:
            # Muchos extractos no dan las participaciones: se calculan con el precio de ese día.
            precio = valor_en(res.serie(ref, fecha), fecha, margen=6)
            if not precio:
                plan.error(n, "No trae unidades y no encuentro el precio de ese día para calcularlas.")
                continue
            neto = importe - (comision or 0) if tipo == "compra" else importe + (comision or 0)
            unidades = neto / precio
            marcas.append("unidades calculadas")
        plan.movimientos.append({"fila": n, "producto": ref, "fecha": fecha, "tipo": tipo,
                                 "unidades": unidades, "importe": round(importe, 2),
                                 "comision": round(comision, 2) if comision else None,
                                 "nota": nota, "marcas": marcas, "origen": "importado"})
    empareja_traspasos(plan, traspasos)
    return plan


def empareja_traspasos(plan, filas):
    """Cada traspaso son dos filas: la salida del fondo de origen y la entrada en el de
    destino, con la misma fecha y (casi) el mismo importe. Aquí se juntan en uno."""
    entradas = [f for f in filas if f["tipo"] == "traspaso_entrada"]
    for s in [f for f in filas if f["tipo"] == "traspaso_salida"]:
        pareja = next((e for e in entradas if e["fecha"] == s["fecha"] and e["producto"] != s["producto"]
                       and abs(e["importe"] - s["importe"]) <= max(1.0, 0.01 * s["importe"])), None)
        if not pareja:
            plan.error(s["fila"], "Traspaso de salida sin su entrada: añade otra fila «traspaso_entrada» con el "
                                  "fondo de destino, la misma fecha y el mismo importe.")
            continue
        entradas.remove(pareja)
        plan.movimientos.append({"fila": s["fila"], "producto": s["producto"], "destino": pareja["producto"],
                                 "fecha": s["fecha"], "tipo": "traspaso", "unidades": s["unidades"],
                                 "unidadesDestino": pareja["unidades"], "importe": s["importe"],
                                 "comision": None, "nota": s["nota"] or pareja["nota"],
                                 "marcas": s["marcas"] + [m for m in pareja["marcas"] if m not in s["marcas"]],
                                 "origen": "importado"})
    for e in entradas:
        plan.error(e["fila"], "Traspaso de entrada sin su salida: añade otra fila «traspaso_salida» con el "
                              "fondo de origen, la misma fecha y el mismo importe.")


def preparar_myinvestor(cfg, archivos, carpeta):
    """[(nombre de archivo, bytes)] de MyInvestor -> Plan. Cada archivo sustituye lo que se
    importó antes de ese fondo, porque el extracto siempre trae la foto completa."""
    plan = Plan()
    res = Resolutor(cfg, plan, carpeta)
    for nombre, contenido in archivos:
        m = re.search(r"([A-Z]{2}[A-Z0-9]{9}\d)", (nombre or "").upper())
        if not m:
            plan.error(nombre, "No encuentro el ISIN en el nombre del archivo. Descárgalo otra vez de "
                               "MyInvestor sin cambiarle el nombre (lleva el ISIN del fondo).")
            continue
        isin = m.group(1)
        lotes, reembolsos = leer_csv_myinvestor(contenido)
        if not lotes and not reembolsos:
            plan.error(nombre, "No parece un extracto «Plusvalías y minusvalías» de MyInvestor, o está vacío.")
            continue
        if len(_coincidencias(cfg, isin, "")) > 1:
            plan.error(nombre, f"«{isin}» está en varios productos (distintas carteras). El CSV de MyInvestor "
                               "no dice a cuál va: impórtalo con la plantilla poniendo la columna «cartera», "
                               "o deja un solo producto con ese ISIN.")
            continue
        ref = res.producto(nombre, isin, "", "fondo")
        if not ref:
            continue
        p = res.datos(ref)
        if p.get("fuente") == "manual":
            plan.error(nombre, f"«{p.get('corto') or p['nombre']}» se valora a mano: el extracto de MyInvestor "
                               "necesita un fondo con precio en internet.")
            continue
        serie = res.serie(ref, lotes[-1][0] if lotes else None)
        plan.series_mi[ref] = serie
        movs, sin_vl, traspasos = myinvestor_a_movimientos(ref, lotes, reembolsos, serie)
        if sin_vl:
            plan.avisos.append(f"{isin}: {sin_vl} compras sin valor liquidativo de ese día; sus "
                               "participaciones se han estimado con el valor del extracto.")
        if traspasos:
            de_traspaso = [m for m in movs if m.get("_traspaso")]
            # Si ya se contestó en una importación anterior, se propone la misma respuesta.
            previo = next((m for m in cfg.get("movimientos", []) if m.get("tipo") == "traspaso"
                           and m.get("destino") == ref and m.get("origen") == "myinvestor"), None)
            plan.traspasos_mi.append({
                "ref": ref, "isin": isin, "lotes": len(de_traspaso),
                "coste": round(sum(m["importe"] for m in de_traspaso), 2),
                "unidades": round(sum(m["unidades"] for m in de_traspaso), 6), "serie": serie,
                "detalle": [[m["fecha"], m["importe"], m["unidades"]] for m in de_traspaso],
                "sugerencia": {"origen": previo["producto"],
                               "fecha": "estimar" if previo.get("fechaEstimada") else previo["fecha"]} if previo else None})
        plan.reemplazar.append((ref, "myinvestor"))
        for mv in movs:
            plan.movimientos.append({**mv, "fila": nombre, "comision": None, "marcas": [],
                                     "origen": "myinvestor"})
    return plan


# ---------------------------------------------------------------- aplicar

def respuesta_valida(item, respuesta, real):
    """La respuesta a «¿de qué fondo venía este traspaso?», si está completa: el fondo de origen y
    una fecha por lote ('fechas', en el orden de item["detalle"]); cada una, una fecha o «estimar».
    Una sola 'fecha' vale para todos los lotes."""
    if not isinstance(respuesta, dict) or not respuesta.get("origen") or respuesta.get("origen") == "ninguno":
        return None
    n = len(item.get("detalle") or []) or 1
    fechas = respuesta.get("fechas")
    if not isinstance(fechas, list):
        fechas = [respuesta.get("fecha")] * n
    fechas = (list(fechas) + [None] * n)[:n]
    out = []
    for f in fechas:
        if f == "estimar":
            out.append("estimar")
            continue
        f = lee_fecha(f)
        if not f:
            return None
        out.append(f)
    return {"origen": respuesta["origen"], "fechas": out, "fecha": out[0]}


def estima_fecha(serie_origen, serie_destino, lote):
    """El día más probable de un traspaso: aquel en que lo que salió del origen (las participaciones
    que se compraron con ese coste en su fecha) vale lo mismo que lo que entró en el destino.
    Puede haber varios días parecidos: es una estimación y se marca como tal. None si no hay precios."""
    fecha, coste, unidades_destino = lote[0], lote[1], lote[2] if len(lote) > 2 else None
    vl_compra = valor_en(serie_origen, fecha, margen=7)
    if not vl_compra or not serie_destino or not unidades_destino:
        return None
    unidades_origen = coste / vl_compra
    mejor, error = None, None
    for d in sorted(serie_destino):
        if d < fecha:
            continue
        vo = valor_en(serie_origen, d, margen=7)
        if not vo:
            continue
        e = abs(unidades_destino * serie_destino[d] / (unidades_origen * vo) - 1)
        if error is None or e < error:
            mejor, error = d, e
    return mejor


def crea_traspaso_mi(cfg, item, resp, destino, serie_origen=None):
    """Convierte los lotes de traspaso de MyInvestor en traspasos desde el fondo de origen: uno
    por cada fecha (cada lote puede tener la suya, o «estimar»). Devuelve ([movimientos], aviso)
    o lanza ErrorValidacion."""
    detalle = item.get("detalle") or []
    fechas = resp.get("fechas") or [resp.get("fecha")] * max(1, len(detalle))
    if not detalle:
        mov, aviso = _un_traspaso(cfg, item, {"origen": resp["origen"], "fecha": fechas[0]}, destino, serie_origen)
        return [mov], aviso
    grupos = {}
    for lote, f in zip(detalle, fechas):
        estimada = f == "estimar"
        if estimada:
            f = estima_fecha(serie_origen or {}, item.get("serie") or {}, lote)
            if not f:
                raise almacen.ErrorValidacion([
                    f"No puedo estimar la fecha del traspaso del lote del {almacen.fmt_fecha(lote[0])} "
                    f"({almacen.fmt_num(lote[1])} €): faltan precios. Pon la fecha a mano."])
        grupos.setdefault((f, estimada), []).append(lote)
    movs, avisos = [], []
    for (f, estimada), lotes in sorted(grupos.items()):
        sub = {**item, "detalle": lotes, "coste": round(sum(l[1] for l in lotes), 2),
               "unidades": sum(l[2] for l in lotes if len(l) > 2)}
        mov, aviso = _un_traspaso(cfg, sub, {"origen": resp["origen"], "fecha": f}, destino, serie_origen)
        if estimada:
            mov["fechaEstimada"] = True
            avisos.append(f"Fecha estimada: {almacen.fmt_fecha(f)}. Es el día en que más cuadran los precios de los dos "
                          "fondos, pero puede no ser el bueno: compruébala en el historial de órdenes de MyInvestor.")
        if aviso:
            avisos.append(aviso)
        movs.append(mov)
    return movs, " ".join(dict.fromkeys(avisos)) or None


def _un_traspaso(cfg, item, resp, destino, serie_origen=None):
    """Un traspaso desde el fondo de origen con los lotes de 'item', en la fecha de 'resp'.
    Devuelve (movimiento, aviso) o lanza ErrorValidacion.

    Si el origen viene de un extracto de MyInvestor, ese extracto solo trae lo que QUEDA en él:
    lo que salió por el traspaso ya no está. Cada lote que salió (con su fecha y su coste, que
    son los del origen) se busca entre sus compras; si no está, se reconstruye: una compra en su
    fecha con las participaciones de ese día. Así el origen sube cuando entró el dinero y baja el
    día del traspaso, y lo que le queda sigue siendo lo del extracto.
    Si el origen se anotó de otra forma, del origen salen, empezando por las más antiguas (FIFO,
    como Hacienda), las participaciones que suman el coste que dice MyInvestor."""
    origen = almacen.producto(cfg, resp["origen"])
    if not origen or origen["id"] == destino:
        raise almacen.ErrorValidacion(["Elige como origen un fondo distinto del de destino."])
    nombre = origen.get("corto") or origen["nombre"]
    del_extracto = [m for m in cfg.get("movimientos", []) if m.get("producto") == origen["id"]
                    and m.get("tipo") == "compra" and m.get("origen") == "myinvestor"]
    if del_extracto and item.get("detalle"):
        return _traspaso_desde_extracto(cfg, item, resp, destino, origen, del_extracto, serie_origen or {})
    lotes = motor.lotes_vivos(cfg, origen["id"], resp["fecha"])
    coste_total = sum(l[1] for l in lotes)
    if not lotes or coste_total <= 0:
        raise almacen.ErrorValidacion([f"En Rumbo, «{nombre}» no tenía nada el {almacen.fmt_fecha(resp['fecha'])}. "
                                       "Si ya lo has reimportado o borrado, elige «No están en Rumbo»."])
    # Las participaciones que salen: las de los lotes más antiguos hasta sumar el coste traspasado.
    falta, unidades = item["coste"], 0.0
    for lu, lc, _, _ in lotes:
        if falta <= 0.005:
            break
        toma = min(lc, falta)
        unidades += lu * toma / lc if lc else lu
        falta -= toma
    aviso = None
    if falta > max(1.0, 0.02 * item["coste"]):
        aviso = (f"«{nombre}» solo tenía {almacen.fmt_num(round(coste_total, 2))} € de coste en Rumbo y el traspaso es de "
                 f"{almacen.fmt_num(item['coste'])} €: sale todo lo que tenía. Revisa que esté completo.")
    precio = valor_en(item.get("serie") or {}, resp["fecha"], margen=7)
    importe = round(item["unidades"] * precio, 2) if precio else item["coste"]
    mov = almacen.guarda_movimiento(cfg, {
        "tipo": "traspaso", "producto": origen["id"], "destino": destino, "fecha": resp["fecha"],
        "unidades": unidades, "unidadesDestino": item["unidades"], "importe": importe,
        "nota": "Traspaso importado de MyInvestor"})
    mov["origen"] = "myinvestor"
    return mov, aviso


def _traspaso_desde_extracto(cfg, item, resp, destino, origen, del_extracto, serie_origen):
    nombre = origen.get("corto") or origen["nombre"]
    usados, unidades, reconstruir = set(), 0.0, []
    for fecha, coste, *_ in item["detalle"]:
        if fecha > resp["fecha"]:
            raise almacen.ErrorValidacion([f"La fecha del traspaso ({almacen.fmt_fecha(resp['fecha'])}) es anterior a la "
                                           f"compra que se traspasó ({almacen.fmt_fecha(fecha)}). Revisa la fecha."])
        par = next((m for m in del_extracto if id(m) not in usados and m["fecha"] == fecha
                    and abs(float(m.get("importe") or 0) - coste) <= max(0.02, coste * 0.001)), None)
        if par:   # el extracto del origen es de antes del traspaso: ese lote sigue ahí
            usados.add(id(par))
            unidades += float(par.get("unidades") or 0)
        else:
            reconstruir.append((fecha, coste))
    for fecha, coste in reconstruir:
        vl = valor_en(serie_origen, fecha, margen=7)
        if not vl:
            raise almacen.ErrorValidacion([f"No encuentro el valor liquidativo de «{nombre}» el {almacen.fmt_fecha(fecha)}: "
                                           "no puedo reconstruir lo que salió de él. Elige «No está en Rumbo»."])
        compra = almacen.guarda_movimiento(cfg, {
            "tipo": "compra", "producto": origen["id"], "fecha": fecha, "unidades": coste / vl, "importe": coste,
            "nota": "Reconstruido: se compró aquí y salió después en un traspaso (el extracto de MyInvestor ya no lo trae)"})
        compra["origen"] = "myinvestor-traspaso"
        compra["paraTraspaso"] = destino
        unidades += coste / vl
    precio = valor_en(item.get("serie") or {}, resp["fecha"], margen=7)
    importe = round(item["unidades"] * precio, 2) if precio else item["coste"]
    mov = almacen.guarda_movimiento(cfg, {
        "tipo": "traspaso", "producto": origen["id"], "destino": destino, "fecha": resp["fecha"],
        "unidades": unidades, "unidadesDestino": item["unidades"], "importe": importe,
        "nota": "Traspaso importado de MyInvestor"})
    mov["origen"] = "myinvestor"
    aviso = (f"El extracto de «{nombre}» ya no trae lo que salió en este traspaso: Rumbo lo ha añadido en su fecha "
             f"({len(reconstruir)} {'compra' if len(reconstruir) == 1 else 'compras'}) para que se vea cómo entró y salió."
             if reconstruir else None)
    return mov, aviso


def aplicar(cfg, plan, respuestas=None, carpeta_series=None):
    """Ejecuta el plan sobre cfg (la real o una copia) y devuelve el informe.
    'respuestas' dice de qué fondo venía cada traspaso de MyInvestor: {ref: {"origen", "fecha"}}.
    'carpeta_series': carpeta de datos, para el precio de un fondo de origen que no viene en este extracto."""
    ids, errores = {}, list(plan.errores)
    for n in plan.productos_nuevos:
        prod, _ = almacen.guarda_producto(cfg, dict(n["datos"]))
        ids[n["ref"]] = prod["id"]
    real = lambda ref: ids.get(ref, ref)

    sustituidos = 0
    for ref, origen in plan.reemplazar:
        antes = len(cfg["movimientos"])
        # También el traspaso que se creó en otra importación hacia este fondo: se vuelve a crear.
        cfg["movimientos"] = [m for m in cfg["movimientos"]
                              if not (m.get("producto") == real(ref) and m.get("origen") == origen)
                              and not (m.get("tipo") == "traspaso" and m.get("destino") == real(ref)
                                       and m.get("origen") == origen)
                              # y lo reconstruido en el fondo de origen para ese traspaso
                              and not (m.get("origen") == "myinvestor-traspaso" and m.get("paraTraspaso") == real(ref))]
        sustituidos += antes - len(cfg["movimientos"])

    respuestas = respuestas or {}
    contestados = {it["ref"]: respuesta_valida(it, respuestas.get(it["ref"]), real) for it in plan.traspasos_mi}

    def firma(m):
        return (m["producto"], m["fecha"], m["tipo"], round(float(m.get("unidades") or 0), 4),
                round(float(m.get("importe") or 0), 2))
    existentes = {firma(m) for m in cfg["movimientos"]}

    añadidos, duplicados, filas = 0, 0, []
    orden = {"compra": 0, "comision": 1, "dividendo": 2, "venta": 3, "traspaso": 4, "ajuste": 6}
    apartados = []   # lotes de traspaso que pasan a ser un traspaso de verdad (si sale bien)
    for mv in sorted(plan.movimientos, key=lambda m: (m["fecha"], orden.get(m["tipo"], 9))):
        if mv.get("_traspaso") and contestados.get(mv["producto"]):
            apartados.append(mv)
            continue
        datos = {k: mv[k] for k in ("fecha", "tipo", "unidades", "importe", "comision", "nota", "unidadesDestino")
                 if mv.get(k) is not None}
        datos["producto"] = real(mv["producto"])
        if mv.get("destino"):
            datos["destino"] = real(mv["destino"])
        estado = "nuevo"
        if firma({**datos, "unidades": datos.get("unidades", 0)}) in existentes:
            duplicados += 1
            estado = "repetido"
        else:
            try:
                guardado = almacen.guarda_movimiento(cfg, datos)
                guardado["origen"] = mv["origen"]
                if mv["tipo"] == "traspaso":
                    existentes.add(firma(guardado))
                añadidos += 1
            except almacen.ErrorValidacion as e:
                errores.append({"fila": mv["fila"], "mensaje": " ".join(e.errores)})
                estado = "error"
        filas.append({"fila": mv["fila"], "fecha": mv["fecha"], "producto": datos["producto"], "tipo": mv["tipo"],
                      "unidades": mv.get("unidades"), "importe": mv["importe"], "marcas": mv["marcas"],
                      "estado": estado})

    # Los traspasos de MyInvestor con respuesta: después de todas las compras (el fondo de
    # origen puede venir en esta misma importación).
    traspasos_mi = []
    for it in plan.traspasos_mi:
        resp = contestados.get(it["ref"])
        estado = {"ref": it["ref"], "isin": it["isin"], "lotes": it["lotes"], "coste": it["coste"],
                  "unidades": it["unidades"], "sugerencia": it["sugerencia"], "detalle": it.get("detalle") or [],
                  "respuesta": respuestas.get(it["ref"]) if isinstance(respuestas.get(it["ref"]), dict) else None,
                  "hecho": False, "error": None, "aviso": None}
        if resp:
            try:
                origen_ref = next((r for r in plan.series_mi if real(r) == resp["origen"]), None)
                serie_origen = plan.series_mi.get(origen_ref) if origen_ref else None
                if serie_origen is None:
                    o = almacen.producto(cfg, resp["origen"])
                    serie_origen = (motor.serie_producto(o, carpeta_series)
                                    if o and o.get("fuente") != "manual" and carpeta_series else {})
                movs, estado["aviso"] = crea_traspaso_mi(cfg, it, resp, real(it["ref"]), serie_origen)
                estado["hecho"] = True
                estado["fechas"] = [m["fecha"] for m in movs]
                for mov in movs:
                    existentes.add(firma(mov))
                    filas.append({"fila": it["isin"], "fecha": mov["fecha"], "producto": real(it["ref"]),
                                  "tipo": "traspaso", "unidades": mov["unidadesDestino"], "importe": mov["importe"],
                                  "marcas": ["traspaso desde " + (almacen.producto(cfg, resp["origen"]) or {}).get("nombre", "")]
                                  + (["fecha estimada"] if mov.get("fechaEstimada") else []),
                                  "estado": "nuevo"})
                    añadidos += 1
            except almacen.ErrorValidacion as e:
                estado["error"] = " ".join(e.errores)
                # No se pierde nada: los lotes se importan como antes (compras con su fecha y coste).
                for mv in [m for m in apartados if m["producto"] == it["ref"]]:
                    datos = {k: mv[k] for k in ("fecha", "tipo", "unidades", "importe", "nota") if mv.get(k) is not None}
                    datos["producto"] = real(mv["producto"])
                    guardado = almacen.guarda_movimiento(cfg, datos)
                    guardado["origen"] = mv["origen"]
                    añadidos += 1
                    filas.append({"fila": mv["fila"], "fecha": mv["fecha"], "producto": datos["producto"], "tipo": "compra",
                                  "unidades": mv.get("unidades"), "importe": mv["importe"], "marcas": ["traspaso"],
                                  "estado": "nuevo"})
        traspasos_mi.append(estado)

    saldos = 0
    for v in plan.valoraciones:
        try:
            almacen.guarda_valoracion(cfg, {"producto": real(v["producto"]), "fecha": v["fecha"], "valor": v["valor"]})
            saldos += 1
            filas.append({"fila": v["fila"], "fecha": v["fecha"], "producto": real(v["producto"]), "tipo": "saldo",
                          "unidades": None, "importe": v["valor"], "marcas": [], "estado": "nuevo"})
        except almacen.ErrorValidacion as e:
            errores.append({"fila": v["fila"], "mensaje": " ".join(e.errores)})

    def total(t):
        return round(sum(f["importe"] for f in filas if f["tipo"] == t and f["estado"] == "nuevo"), 2)
    nombres = {p["id"]: p.get("corto") or p["nombre"] for p in cfg["productos"]}
    for f in filas:
        f["productoNombre"] = nombres.get(f["producto"], f["producto"])
    for t in traspasos_mi:
        t["nombre"] = nombres.get(real(t["ref"]), t["isin"])
        # Posibles fondos de origen: todo lo que no es una cuenta, menos el propio destino.
        t["origenes"] = [{"id": p["id"], "nombre": p.get("corto") or p["nombre"]} for p in cfg["productos"]
                         if p["id"] != real(t["ref"]) and p.get("tipo") not in almacen.SOLO_SALDO]
    return {
        "añadidos": añadidos, "repetidos": duplicados, "saldos": saldos, "sustituidos": sustituidos,
        "errores": sorted(errores, key=lambda e: (str(type(e["fila"])), str(e["fila"]).zfill(6))),
        "avisos": plan.avisos,
        "totales": {"compras": total("compra"), "ventas": total("venta"),
                    "dividendos": total("dividendo"), "comisiones": total("comision"),
                    "numCompras": sum(1 for f in filas if f["tipo"] == "compra" and f["estado"] == "nuevo")},
        "productosNuevos": [{**n["datos"], "id": ids.get(n["ref"]), "precio": n["precio"],
                             "monedaPrecio": n["monedaPrecio"], "fechaPrecio": n["fecha"],
                             "mercado": n["mercado"]} for n in plan.productos_nuevos],
        "filas": sorted(filas, key=lambda f: f["fecha"], reverse=True),
        "traspasosMyInvestor": traspasos_mi,
    }


def vista_previa(cfg, plan, respuestas=None, carpeta_series=None):
    """El mismo informe que dará aplicar(), sin tocar la cartera de verdad."""
    return aplicar(copy.deepcopy(cfg), plan, respuestas, carpeta_series)
